from pathlib import Path
import os
from collections import defaultdict
from docx import Document
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from app.core.config import get_settings
from app.models.entities import TestRun


def generate_reports(test: TestRun) -> tuple[str, str]:
    output = Path("/tmp/reports/generated") if os.getenv("VERCEL") == "1" else Path("reports/generated")
    output.mkdir(parents=True, exist_ok=True)
    pdf_path = output / f"{test.id}.pdf"
    docx_path = output / f"{test.id}.docx"

    findings = test.details.get("findings", []) if isinstance(test.details, dict) else []
    category_summary = _build_category_summary(findings)

    pdf = canvas.Canvas(str(pdf_path), pagesize=letter)
    pdf.setTitle(f"AI Red-Team Report {test.id}")
    pdf.drawString(72, 740, "Sentinel Red AI - Guardrail Test Report")
    pdf.drawString(72, 715, f"Model: {test.target_model}")
    pdf.drawString(72, 695, f"Risk Score: {test.risk_score}/100")
    pdf.drawString(72, 675, f"Attack Success Rate: {test.attack_success_rate}%")
    pdf.drawString(72, 655, f"Hallucination Rate: {test.hallucination_rate}%")
    pdf.drawString(72, 635, f"Recommendation: {test.recommendation[:100]}")

    pdf.drawString(72, 605, "Category summary:")
    y = 585
    pdf.setFont("Helvetica", 9)
    pdf.drawString(72, y, "Category")
    pdf.drawString(190, y, "Avg Risk")
    pdf.drawString(260, y, "Attack Success")
    pdf.drawString(360, y, "Findings")
    y -= 14

    for category in category_summary:
        pdf.drawString(72, y, category["category"].replace("_", " ").title())
        pdf.drawString(190, y, f"{category['avg_risk']}%")
        pdf.drawString(260, y, f"{category['attack_success_rate']}%")
        pdf.drawString(360, y, str(category["count"]))
        y -= 14

    y -= 20
    pdf.drawString(72, y, "Risk profile chart:")
    y -= 14
    chart_x = 72
    chart_y = y - 100
    chart_width = 420
    bar_height = 14

    max_bar_value = max((category["avg_risk"] for category in category_summary), default=100)
    for item in category_summary:
        pdf.setFillColor(colors.grey)
        pdf.rect(chart_x, chart_y, chart_width, bar_height, fill=0, stroke=1)
        bar_length = (item["avg_risk"] / max_bar_value) * chart_width if max_bar_value else 0
        pdf.setFillColor(colors.HexColor("#4F46E5"))
        pdf.rect(chart_x, chart_y, bar_length, bar_height, fill=1, stroke=0)
        pdf.setFillColor(colors.black)
        pdf.drawString(chart_x + chart_width + 10, chart_y + 2, f"{item['category'].replace('_', ' ').title()} {item['avg_risk']}%")
        chart_y -= bar_height + 10

    pdf.setFont("Helvetica", 10)
    pdf.drawString(72, chart_y - 20, "Top unique recommendations and findings:")
    y = chart_y - 40
    for finding in _unique_category_findings(findings)[:6]:
        pdf.drawString(72, y, f"- {finding}")
        y -= 14

    pdf.save()

    doc = Document()
    doc.add_heading("Sentinel Red AI - Guardrail Test Report", level=1)
    doc.add_paragraph(f"Model: {test.target_model}")
    doc.add_paragraph(f"Risk Score: {test.risk_score}/100")
    doc.add_paragraph(f"Attack Success Rate: {test.attack_success_rate}%")
    doc.add_paragraph(f"Hallucination Rate: {test.hallucination_rate}%")
    doc.add_heading("Defense Recommendation", level=2)
    doc.add_paragraph(test.recommendation)
    doc.add_heading("Category Summary", level=2)

    table = doc.add_table(rows=1, cols=4)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Category"
    hdr_cells[1].text = "Avg Risk"
    hdr_cells[2].text = "Attack Success"
    hdr_cells[3].text = "Findings"
    for category in category_summary:
        row_cells = table.add_row().cells
        row_cells[0].text = category["category"].replace("_", " ").title()
        row_cells[1].text = f"{category['avg_risk']}%"
        row_cells[2].text = f"{category['attack_success_rate']}%"
        row_cells[3].text = str(category["count"])

    doc.add_heading("Top Findings", level=2)
    for finding in _unique_category_findings(findings)[:8]:
        doc.add_paragraph(finding, style="List Bullet")
    doc.save(str(docx_path))
    return str(pdf_path), str(docx_path)


def _build_category_summary(findings: list[dict]) -> list[dict]:
    buckets = defaultdict(lambda: {"risk": 0, "count": 0, "failures": 0})
    for finding in findings:
        category = finding.get("category", "unknown")
        buckets[category]["risk"] += finding.get("risk", 0)
        buckets[category]["count"] += 1
        if not finding.get("blocked", True):
            buckets[category]["failures"] += 1

    summary = []
    for category, stats in buckets.items():
        count = stats["count"]
        avg_risk = round(stats["risk"] / max(count, 1), 2)
        attack_success_rate = round((stats["failures"] / max(count, 1)) * 100, 2)
        summary.append({
            "category": category,
            "count": count,
            "avg_risk": avg_risk,
            "attack_success_rate": attack_success_rate,
        })
    summary.sort(key=lambda item: item["avg_risk"], reverse=True)
    return summary


def _unique_category_findings(findings: list[dict]) -> list[str]:
    seen = set()
    unique = []
    for finding in findings:
        text = f"{finding.get('category', 'unknown').replace('_', ' ').title()}: {finding.get('finding', '')}"
        if text in seen:
            continue
        seen.add(text)
        unique.append(text)
    return unique


def public_url(path: str) -> str:
    settings = get_settings()
    return f"{settings.report_base_url}/static/{Path(path).name}"
